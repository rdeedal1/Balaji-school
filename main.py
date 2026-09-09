import customtkinter as ctk
import sqlite3
from tkinter import messagebox, Toplevel, filedialog, ttk
from tkcalendar import Calendar
from datetime import datetime
import calendar
import re
import openpyxl

ctk.set_appearance_mode("System")
ctk.set_default_color_theme("blue")

# --- डेटाबेस सेटअप ---
def init_db():
    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS students (
            sr_no TEXT PRIMARY KEY,
            student_name TEXT,
            father_name TEXT,
            mother_name TEXT,
            dob TEXT,
            gender TEXT,
            category TEXT,
            address TEXT,
            doa TEXT,
            rte TEXT,
            class_name TEXT,
            mobile TEXT,
            route TEXT
        )
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS attendance (
            date TEXT,
            sr_no TEXT,
            class_name TEXT,
            status TEXT,
            PRIMARY KEY (date, sr_no)
        )
    """)
    conn.commit()
    conn.close()

init_db()

route_options = [
    "Self", "Kheerwa", "Kheri", "Pachkodiya", "Bhojpura", "Mundoti", 
    "Thikariya", "sunderpura", "Ghar", "Kakraliyo ki dhani", 
    "Tetarwalo ki dhani", "Sukalpura", "Bheslana", "Minda"
]

months_list = ["01 - Jan", "02 - Feb", "03 - Mar", "04 - Apr", "05 - May", "06 - Jun", 
               "07 - Jul", "08 - Aug", "09 - Sep", "10 - Oct", "11 - Nov", "12 - Dec"]
years_list = [str(y) for y in range(2024, 2030)]

CLASS_ORDER = ['PP.3+', 'PP.4+', 'PP.5+', '1', '2', '3', '4', '5', '6', '7', '8', '9', '10']

def normalize_class_name(c_name):
    c = str(c_name).strip().upper().replace(" ", "").replace("-", "")
    if "PP3" in c or "PP.3" in c or "NUR" in c:
        return "PP.3+"
    if "PP4" in c or "PP.4" in c or "LKG" in c or "KG1" in c:
        return "PP.4+"
    if "PP5" in c or "PP.5" in c or "UKG" in c or "KG2" in c:
        return "PP.5+"
    
    digits = re.findall(r'\d+', c)
    if digits:
        val = str(int(digits[0]))
        if val in CLASS_ORDER:
            return val
    return str(c_name).strip()

def sort_classes(classes):
    seen_normalized = {}
    for c in classes:
        if str(c).strip():
            norm = normalize_class_name(c)
            if norm not in seen_normalized:
                seen_normalized[norm] = norm

    def get_sort_index(norm_val):
        if norm_val in CLASS_ORDER:
            return (0, CLASS_ORDER.index(norm_val))
        digits = re.findall(r'\d+', norm_val)
        if digits:
            return (1, int(digits[0]))
        return (2, str(norm_val))

    sorted_list = sorted(seen_normalized.keys(), key=get_sort_index)
    return sorted_list

# --- कैलेंडर पॉप-अप ---
def open_calendar(target_entry, on_select_callback=None):
    top = Toplevel(app)
    top.title("तारीख चुनें")
    top.geometry("280x260")
    top.resizable(False, False)
    top.grab_set()

    today = datetime.today()
    cal = Calendar(top, selectmode='day', year=today.year, month=today.month, day=today.day, date_pattern='dd-mm-yyyy')
    cal.pack(pady=10)

    def set_date():
        target_entry.delete(0, "end")
        target_entry.insert(0, cal.get_date())
        top.destroy()
        if on_select_callback:
            on_select_callback()

    ctk.CTkButton(top, text="Select", width=100, command=set_date).pack(pady=5)

# --- डेटा सेव या अपडेट ---
def save_or_update():
    sr_no = entry_sr.get().strip()
    name = entry_name.get().strip()
    father = entry_father.get().strip()
    mother = entry_mother.get().strip()
    dob = entry_dob.get().strip()
    gender = opt_gender.get()
    category = opt_category.get()
    address = entry_address.get().strip()
    doa = entry_doa.get().strip()
    rte = opt_rte.get()
    cls = normalize_class_name(entry_class.get().strip())
    mobile = entry_mobile.get().strip()
    route = opt_route.get()

    if not sr_no or not name or not cls:
        messagebox.showerror("Error", "SR No., Student Name और Class भरना अनिवार्य है!")
        return

    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()

    if is_edit_mode.get():
        cursor.execute("""
            UPDATE students SET 
                student_name = ?, father_name = ?, mother_name = ?, 
                dob = ?, gender = ?, category = ?, address = ?, 
                doa = ?, rte = ?, class_name = ?, mobile = ?, route = ?
            WHERE sr_no = ?
        """, (name, father, mother, dob, gender, category, address, doa, rte, cls, mobile, route, sr_no))
        conn.commit()
        conn.close()
        messagebox.showinfo("Success", f"SR No. {sr_no} ({name}) का डेटा सफलतापूर्वक अपडेट हो गया!")
        clear_fields()
        load_records()
        update_class_dropdown()
        refresh_dashboard()
        tabview.set("📋 Student Records")
    else:
        try:
            cursor.execute("INSERT INTO students VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                           (sr_no, name, father, mother, dob, gender, category, address, doa, rte, cls, mobile, route))
            conn.commit()
            conn.close()
            messagebox.showinfo("Success", f"SR No. {sr_no} ({name}) का डेटा सुरक्षित हो गया!")
            clear_fields()
            load_records()
            update_class_dropdown()
            refresh_dashboard()
        except sqlite3.IntegrityError:
            conn.close()
            messagebox.showerror("Error", f"SR No. {sr_no} पहले से दर्ज है!")
        except Exception as e:
            conn.close()
            messagebox.showerror("Error", f"त्रुटि: {str(e)}")

# --- एक्सेल इम्पोर्ट फंक्शन ---
def import_excel():
    file_path = filedialog.askopenfilename(title="एक्सेल फ़ाइल चुनें", filetypes=[("Excel Files", "*.xlsx *.xls")])
    if not file_path:
        return

    try:
        wb = openpyxl.load_workbook(file_path)
        sheet = wb.active
        success_count = 0
        skip_count = 0

        conn = sqlite3.connect("students.db")
        cursor = conn.cursor()

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue

            sr_no = str(row[0]).strip() if row[0] is not None else ""
            name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            father = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            mother = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
            
            dob = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
            if isinstance(row[4], datetime):
                dob = row[4].strftime("%d-%m-%Y")

            gender = str(row[5]).strip() if len(row) > 5 and row[5] is not None else "Male"
            category = str(row[6]).strip() if len(row) > 6 and row[6] is not None else "General"
            address = str(row[7]).strip() if len(row) > 7 and row[7] is not None else ""

            doa = str(row[8]).strip() if len(row) > 8 and row[8] is not None else ""
            if isinstance(row[8], datetime):
                doa = row[8].strftime("%d-%m-%Y")

            rte = str(row[9]).strip() if len(row) > 9 and row[9] is not None else "No"
            raw_cls = str(row[10]).strip() if len(row) > 10 and row[10] is not None else ""
            cls = normalize_class_name(raw_cls)
            mobile = str(row[11]).strip() if len(row) > 11 and row[11] is not None else ""
            route = str(row[12]).strip() if len(row) > 12 and row[12] is not None else "Self"

            if not sr_no or not name:
                skip_count += 1
                continue

            try:
                cursor.execute("INSERT INTO students VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                               (sr_no, name, father, mother, dob, gender, category, address, doa, rte, cls, mobile, route))
                success_count += 1
            except sqlite3.IntegrityError:
                skip_count += 1

        conn.commit()
        conn.close()
        load_records()
        update_class_dropdown()
        refresh_dashboard()
        messagebox.showinfo("Import Complete", f"जुड़े रिकॉर्ड: {success_count}\nछोड़े गए रिकॉर्ड: {skip_count}")

    except Exception as e:
        messagebox.showerror("Error", f"एक्सेल लोड करने में समस्या:\n{str(e)}")

# --- फ़ॉर्म रीसेट ---
def clear_fields():
    is_edit_mode.set(False)
    entry_sr.configure(state="normal")
    entry_sr.delete(0, "end")
    entry_name.delete(0, "end")
    entry_father.delete(0, "end")
    entry_mother.delete(0, "end")
    entry_dob.delete(0, "end")
    entry_doa.delete(0, "end")
    opt_gender.set("Male")
    opt_category.set("General")
    entry_address.delete(0, "end")
    opt_rte.set("No")
    entry_class.delete(0, "end")
    entry_mobile.delete(0, "end")
    opt_route.set("Self")
    btn_save.configure(text="Save Student", fg_color="#1f6aa5")
    lbl_form_title.configure(text="विद्यार्थी प्रवेश फ़ॉर्म (Student Admission Form)")

# --- टेबल डबल क्लिक एडिट ---
def on_double_click_row(event):
    selected = tree.selection()
    if not selected:
        return

    item = tree.item(selected[0])
    sr_no = str(item['values'][0])

    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM students WHERE sr_no = ?", (sr_no,))
    row = cursor.fetchone()
    conn.close()

    if not row:
        return

    clear_fields()
    is_edit_mode.set(True)

    entry_sr.insert(0, str(row[0]))
    entry_sr.configure(state="disabled")
    entry_name.insert(0, str(row[1]) if row[1] else "")
    entry_father.insert(0, str(row[2]) if row[2] else "")
    entry_mother.insert(0, str(row[3]) if row[3] else "")
    entry_dob.insert(0, str(row[4]) if row[4] else "")
    opt_gender.set(str(row[5]) if row[5] else "Male")
    opt_category.set(str(row[6]) if row[6] else "General")
    entry_address.insert(0, str(row[7]) if row[7] else "")
    entry_doa.insert(0, str(row[8]) if row[8] else "")
    opt_rte.set(str(row[9]) if row[9] else "No")
    entry_class.insert(0, str(row[10]) if row[10] else "")
    entry_mobile.insert(0, str(row[11]) if row[11] else "")
    opt_route.set(str(row[12]) if row[12] in route_options else "Self")

    btn_save.configure(text="💾 Update Record", fg_color="#e65100")
    lbl_form_title.configure(text=f"रिकॉर्ड सुधारें (Editing SR No: {sr_no})")
    tabview.set("➕ New Admission")

# --- टेबल लोड / सर्च ---
def load_records(search_query=""):
    for item in tree.get_children():
        tree.delete(item)

    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()

    if search_query:
        wildcard = f"%{search_query}%"
        cursor.execute("""
            SELECT sr_no, student_name, class_name, father_name, mobile, route, category, rte, dob, doa, address 
            FROM students 
            WHERE sr_no LIKE ? OR student_name LIKE ? OR class_name LIKE ? OR route LIKE ?
        """, (wildcard, wildcard, wildcard, wildcard))
    else:
        cursor.execute("SELECT sr_no, student_name, class_name, father_name, mobile, route, category, rte, dob, doa, address FROM students")

    rows = cursor.fetchall()
    for row in rows:
        tree.insert("", "end", values=row)

    conn.close()
    lbl_count.configure(text=f"कुल रिकॉर्ड: {len(rows)}")

def search_data():
    q = entry_search.get().strip()
    load_records(q)

def reset_search():
    entry_search.delete(0, "end")
    load_records()

# --- रिकॉर्ड हटाना ---
def delete_selected():
    selected = tree.selection()
    if not selected:
        messagebox.showwarning("चेतावनी", "कृपया टेबल में से किसी छात्र को चुनें!")
        return

    item = tree.item(selected[0])
    sr_no = str(item['values'][0])
    name = str(item['values'][1])

    confirm = messagebox.askyesno("पुष्टि करें", f"क्या आप SR No. {sr_no} ({name}) को हटाना चाहते हैं?")
    if confirm:
        conn = sqlite3.connect("students.db")
        cursor = conn.cursor()
        cursor.execute("DELETE FROM students WHERE sr_no = ?", (sr_no,))
        cursor.execute("DELETE FROM attendance WHERE sr_no = ?", (sr_no,))
        conn.commit()
        conn.close()
        load_records()
        update_class_dropdown()
        refresh_dashboard()
        messagebox.showinfo("हटा दिया गया", "रिकॉर्ड हटा दिया गया है।")

# --- छात्र डेटा एक्सेल एक्सपोर्ट ---
def export_excel():
    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM students")
    rows = cursor.fetchall()
    conn.close()

    if not rows:
        messagebox.showwarning("खाली डेटा", "एक्सपोर्ट करने के लिए कोई रिकॉर्ड नहीं है!")
        return

    file_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")], title="डेटा सेव करें")
    if not file_path:
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = ["SR No", "Student Name", "Father Name", "Mother Name", "DOB", "Gender", "Category", "Address", "DOA", "RTE", "Class", "Mobile", "Route"]
    ws.append(headers)

    for r in rows:
        ws.append(r)

    wb.save(file_path)
    messagebox.showinfo("Export Success", f"डेटा सफलतापूर्वक सेव कर दिया गया:\n{file_path}")

# ================= ATTENDANCE FUNCTIONS =================
attendance_vars = {}

def update_class_dropdown():
    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()
    cursor.execute("SELECT DISTINCT class_name FROM students WHERE class_name != ''")
    raw_classes = [r[0] for r in cursor.fetchall()]
    conn.close()

    sorted_classes = sort_classes(raw_classes)

    if not sorted_classes:
        sorted_classes = CLASS_ORDER
    opt_att_class.configure(values=sorted_classes)
    opt_att_class.set(sorted_classes[0])
    opt_rep_class.configure(values=sorted_classes)
    opt_rep_class.set(sorted_classes[0])

def set_all_status(status_value):
    if not attendance_vars:
        messagebox.showwarning("चेतावनी", "कृपया पहले 'Load Students' पर क्लिक करके छात्र लोड करें!")
        return
    for var in attendance_vars.values():
        var.set(status_value)

def load_attendance_list():
    selected_class = opt_att_class.get()
    norm_c = normalize_class_name(selected_class)
    att_date = entry_att_date.get().strip()

    if not att_date:
        messagebox.showerror("Error", "कृपया तारीख दर्ज करें!")
        return

    for widget in att_scroll_frame.winfo_children():
        widget.destroy()
    attendance_vars.clear()

    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT sr_no, student_name, father_name 
        FROM students 
        WHERE class_name = ? OR class_name = ?
        ORDER BY sr_no
    """, (selected_class, norm_c))
    students = cursor.fetchall()

    if not students:
        ctk.CTkLabel(att_scroll_frame, text="इस कक्षा में कोई विद्यार्थी दर्ज नहीं है।", font=ctk.CTkFont(size=14)).pack(pady=20)
        conn.close()
        return

    cursor.execute("""
        SELECT sr_no, status FROM attendance 
        WHERE date = ? AND (class_name = ? OR class_name = ?)
    """, (att_date, selected_class, norm_c))
    saved_status = dict(cursor.fetchall())
    conn.close()

    if saved_status:
        lbl_att_status_mode.configure(
            text=f"📌 {att_date} की उपस्थिति पहले से दर्ज है (आप इसे Edit / Update कर सकते हैं)",
            text_color="#ef6c00"
        )
    else:
        lbl_att_status_mode.configure(
            text=f"🆕 {att_date} के लिए नई हाजिरी दर्ज की जा रही है",
            text_color="#2e7d32"
        )

    header_f = ctk.CTkFrame(att_scroll_frame, fg_color=("gray80", "gray25"), height=35)
    header_f.pack(fill="x", padx=5, pady=(2, 6))
    header_f.grid_columnconfigure((0, 1, 2, 3), weight=1)

    ctk.CTkLabel(header_f, text="SR No.", font=ctk.CTkFont(size=13, weight="bold")).grid(row=0, column=0, padx=15, pady=5, sticky="w")
    ctk.CTkLabel(header_f, text="विद्यार्थी का नाम", font=ctk.CTkFont(size=13, weight="bold")).grid(row=0, column=1, padx=10, pady=5, sticky="w")
    ctk.CTkLabel(header_f, text="पिता का नाम", font=ctk.CTkFont(size=13, weight="bold")).grid(row=0, column=2, padx=10, pady=5, sticky="w")
    ctk.CTkLabel(header_f, text="उपस्थिति (Status)", font=ctk.CTkFont(size=13, weight="bold")).grid(row=0, column=3, padx=15, pady=5, sticky="e")

    for row_idx, s in enumerate(students):
        sr = str(s[0])
        name = str(s[1])
        father = str(s[2]) if s[2] else "-"

        current_val = saved_status.get(sr, "Present")
        status_var = ctk.StringVar(value=current_val)
        attendance_vars[sr] = status_var

        row_bg = ("#f0f0f0", "#1f1f1f") if row_idx % 2 == 0 else ("#ffffff", "#2b2b2b")
        row_f = ctk.CTkFrame(att_scroll_frame, fg_color=row_bg, corner_radius=6)
        row_f.pack(fill="x", padx=5, pady=2)
        row_f.grid_columnconfigure((0, 1, 2, 3), weight=1)

        ctk.CTkLabel(row_f, text=sr, font=ctk.CTkFont(size=12)).grid(row=0, column=0, padx=15, pady=6, sticky="w")
        ctk.CTkLabel(row_f, text=name, font=ctk.CTkFont(size=13, weight="bold")).grid(row=0, column=1, padx=10, pady=6, sticky="w")
        ctk.CTkLabel(row_f, text=father, font=ctk.CTkFont(size=12)).grid(row=0, column=2, padx=10, pady=6, sticky="w")

        rb_frame = ctk.CTkFrame(row_f, fg_color="transparent")
        rb_frame.grid(row=0, column=3, padx=10, pady=4, sticky="e")

        ctk.CTkRadioButton(rb_frame, text="P", value="Present", variable=status_var, width=48, fg_color="#2e7d32", text_color=("black", "white")).pack(side="left", padx=1)
        ctk.CTkRadioButton(rb_frame, text="A", value="Absent", variable=status_var, width=48, fg_color="#c62828", text_color=("black", "white")).pack(side="left", padx=1)
        ctk.CTkRadioButton(rb_frame, text="L", value="Leave", variable=status_var, width=48, fg_color="#ef6c00", text_color=("black", "white")).pack(side="left", padx=1)
        ctk.CTkRadioButton(rb_frame, text="H", value="Holiday", variable=status_var, width=48, fg_color="#6a1b9a", text_color=("black", "white")).pack(side="left", padx=1)

    lbl_att_summary.configure(text=f"कुल छात्र: {len(students)} | कक्षा: {norm_c}")

def save_attendance():
    if not attendance_vars:
        messagebox.showwarning("चेतावनी", "सेव करने के लिए कोई छात्र लोड नहीं किया गया है!")
        return

    att_date = entry_att_date.get().strip()
    selected_class = opt_att_class.get()
    norm_c = normalize_class_name(selected_class)

    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()

    present_c = 0
    absent_c = 0
    leave_c = 0
    holiday_c = 0

    for sr, var in attendance_vars.items():
        st = var.get()
        if st == "Present":
            present_c += 1
        elif st == "Absent":
            absent_c += 1
        elif st == "Leave":
            leave_c += 1
        else:
            holiday_c += 1

        cursor.execute("""
            INSERT INTO attendance (date, sr_no, class_name, status) 
            VALUES (?, ?, ?, ?)
            ON CONFLICT(date, sr_no) DO UPDATE SET status = excluded.status
        """, (att_date, sr, norm_c, st))

    conn.commit()
    conn.close()

    summary_text = f"तारीख: {att_date} (Class: {norm_c})\n\nकुल: {len(attendance_vars)}\n"
    if holiday_c == len(attendance_vars):
        summary_text += "🏖️ पूरे दिन की छुट्टी (Holiday) दर्ज की गई!"
    else:
        summary_text += f"🟢 Present: {present_c}\n🔴 Absent: {absent_c}\n🟡 Leave: {leave_c}\n🟣 Holiday: {holiday_c}\n\nडेटा सुरक्षित (Updated) हो गया!"

    refresh_dashboard()
    messagebox.showinfo("Attendance Saved", summary_text)

# ================= ATTENDANCE REPORT FUNCTIONS =================
def load_attendance_report():
    for item in tree_report.get_children():
        tree_report.delete(item)

    selected_class = opt_rep_class.get()
    norm_c = normalize_class_name(selected_class)
    month_val = opt_rep_month.get().split(" - ")[0]
    year_val = opt_rep_year.get()
    date_filter = f"%-{month_val}-{year_val}"

    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT sr_no, student_name, father_name 
        FROM students 
        WHERE class_name = ? OR class_name = ?
        ORDER BY sr_no
    """, (selected_class, norm_c))
    students = cursor.fetchall()

    if not students:
        conn.close()
        lbl_rep_summary.configure(text="इस कक्षा में कोई छात्र नहीं मिले।")
        return

    cursor.execute("""
        SELECT COUNT(DISTINCT date) FROM attendance 
        WHERE (class_name = ? OR class_name = ?) AND date LIKE ? AND status != 'Holiday'
    """, (selected_class, norm_c, date_filter))
    working_days = cursor.fetchone()[0] or 0

    for s in students:
        sr = str(s[0])
        name = str(s[1])
        father = str(s[2]) if s[2] else "-"

        cursor.execute("""
            SELECT 
                SUM(CASE WHEN status = 'Present' THEN 1 ELSE 0 END),
                SUM(CASE WHEN status = 'Absent' THEN 1 ELSE 0 END),
                SUM(CASE WHEN status = 'Leave' THEN 1 ELSE 0 END),
                SUM(CASE WHEN status = 'Holiday' THEN 1 ELSE 0 END)
            FROM attendance 
            WHERE sr_no = ? AND date LIKE ?
        """, (sr, date_filter))
        att_data = cursor.fetchone()

        p_count = att_data[0] if att_data[0] is not None else 0
        a_count = att_data[1] if att_data[1] is not None else 0
        l_count = att_data[2] if att_data[2] is not None else 0
        h_count = att_data[3] if att_data[3] is not None else 0

        pct = f"{(p_count / working_days * 100):.1f}%" if working_days > 0 else "0.0%"

        tree_report.insert("", "end", values=(sr, name, father, norm_c, working_days, p_count, a_count, l_count, h_count, pct))

    conn.close()
    lbl_rep_summary.configure(text=f"कक्षा: {norm_c} | कुल कार्यदिवस (Working Days): {working_days} | कुल विद्यार्थी: {len(students)}")

def export_attendance_monthly():
    selected_class = opt_rep_class.get()
    norm_c = normalize_class_name(selected_class)
    month_str = opt_rep_month.get()
    month_val = int(month_str.split(" - ")[0])
    year_val = int(opt_rep_year.get())

    num_days = calendar.monthrange(year_val, month_val)[1]

    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT sr_no, student_name, father_name 
        FROM students 
        WHERE class_name = ? OR class_name = ?
        ORDER BY sr_no
    """, (selected_class, norm_c))
    students = cursor.fetchall()

    if not students:
        conn.close()
        messagebox.showwarning("खाली डेटा", "एक्सपोर्ट करने के लिए कोई छात्र नहीं मिले!")
        return

    clean_file_class = norm_c.replace(".", "").replace("+", "plus")
    file_path = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel Files", "*.xlsx")],
        initialfile=f"Attendance_Register_{clean_file_class}_{month_val}_{year_val}.xlsx",
        title="मासिक अटेंडेंस रजिस्टर सेव करें"
    )
    if not file_path:
        conn.close()
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Class {clean_file_class}"

    headers = ["SR No", "Student Name", "Father Name"] + [f"{d:02d}" for d in range(1, num_days + 1)] + ["Total P", "Total A", "Total L", "Total H", "%"]
    ws.append(headers)

    for s in students:
        sr = str(s[0])
        name = str(s[1])
        father = str(s[2]) if s[2] else "-"

        row_data = [sr, name, father]
        p_c = 0
        a_c = 0
        l_c = 0
        h_c = 0
        effective_days = 0

        for d in range(1, num_days + 1):
            date_str = f"{d:02d}-{month_val:02d}-{year_val}"
            cursor.execute("SELECT status FROM attendance WHERE sr_no = ? AND date = ?", (sr, date_str))
            res = cursor.fetchone()
            if res:
                st = res[0]
                if st == "Present":
                    row_data.append("P")
                    p_c += 1
                    effective_days += 1
                elif st == "Absent":
                    row_data.append("A")
                    a_c += 1
                    effective_days += 1
                elif st == "Leave":
                    row_data.append("L")
                    l_c += 1
                    effective_days += 1
                elif st == "Holiday":
                    row_data.append("H")
                    h_c += 1
            else:
                row_data.append("-")

        pct = f"{(p_c / effective_days * 100):.1f}%" if effective_days > 0 else "0%"
        row_data.extend([p_c, a_c, l_c, h_c, pct])
        ws.append(row_data)

    conn.close()
    wb.save(file_path)
    messagebox.showinfo("Export Success", f"मासिक अटेंडेंस शीट सफलतापूर्वक सेव हो गई:\n{file_path}")

# ================= ERP DASHBOARD REFRESH =================
def refresh_dashboard():
    selected_date = entry_dash_date.get().strip()
    if not selected_date:
        selected_date = datetime.today().strftime("%d-%m-%Y")
        entry_dash_date.delete(0, "end")
        entry_dash_date.insert(0, selected_date)

    conn = sqlite3.connect("students.db")
    cursor = conn.cursor()

    cursor.execute("SELECT COUNT(*) FROM students")
    total_students = cursor.fetchone()[0] or 0

    cursor.execute("""
        SELECT 
            SUM(CASE WHEN status = 'Present' THEN 1 ELSE 0 END),
            SUM(CASE WHEN status = 'Absent' THEN 1 ELSE 0 END),
            SUM(CASE WHEN status IN ('Leave', 'Holiday') THEN 1 ELSE 0 END)
        FROM attendance WHERE date = ?
    """, (selected_date,))
    att_row = cursor.fetchone()

    p_today = att_row[0] if att_row and att_row[0] is not None else 0
    a_today = att_row[1] if att_row and att_row[1] is not None else 0
    l_today = att_row[2] if att_row and att_row[2] is not None else 0

    cursor.execute("SELECT COUNT(DISTINCT route) FROM students WHERE route != '' AND route != 'Self'")
    active_routes = cursor.fetchone()[0] or 0

    lbl_val_students.configure(text=str(total_students))
    lbl_val_present.configure(text=str(p_today))
    lbl_val_absent.configure(text=str(a_today))
    lbl_val_leave.configure(text=str(l_today))
    lbl_val_routes.configure(text=str(active_routes))

    for item in tree_dash.get_children():
        tree_dash.delete(item)

    cursor.execute("SELECT DISTINCT class_name FROM students WHERE class_name != ''")
    raw_classes = [r[0] for r in cursor.fetchall()]
    
    sorted_classes = sort_classes(raw_classes)

    for norm_c in sorted_classes:
        cursor.execute("SELECT COUNT(*) FROM students WHERE class_name = ? OR class_name = ?", (norm_c, norm_c.replace(".", "")))
        cls_total = cursor.fetchone()[0] or 0

        cursor.execute("""
            SELECT 
                SUM(CASE WHEN status = 'Present' THEN 1 ELSE 0 END),
                SUM(CASE WHEN status = 'Absent' THEN 1 ELSE 0 END),
                SUM(CASE WHEN status IN ('Leave', 'Holiday') THEN 1 ELSE 0 END)
            FROM attendance 
            WHERE date = ? AND (class_name = ? OR class_name = ?)
        """, (selected_date, norm_c, norm_c.replace(".", "")))
        cls_att = cursor.fetchone()

        c_p = cls_att[0] if cls_att and cls_att[0] is not None else 0
        c_a = cls_att[1] if cls_att and cls_att[1] is not None else 0
        c_l = cls_att[2] if cls_att and cls_att[2] is not None else 0

        marked = c_p + c_a + c_l
        if marked > 0:
            status_text = "दर्ज (Marked)"
            pct = f"{(c_p / (c_p + c_a) * 100):.1f}%" if (c_p + c_a) > 0 else "0%"
        else:
            status_text = "लंबित (Pending)"
            pct = "N/A"

        tree_dash.insert("", "end", values=(norm_c, cls_total, c_p, c_a, c_l, pct, status_text))

    conn.close()

# ================= UI विंडो =================
app = ctk.CTk()
app.geometry("1100x780")
app.title("स्कूल स्टूडेंट मैनेजमेंट व ERP सिस्टम")

is_edit_mode = ctk.BooleanVar(value=False)

tabview = ctk.CTkTabview(app)
tabview.pack(fill="both", expand=True, padx=15, pady=10)

tab_dash = tabview.add("📊 ERP Dashboard")
tab_entry = tabview.add("➕ New Admission")
tab_records = tabview.add("📋 Student Records")
tab_attendance = tabview.add("📅 Daily Attendance")
tab_report = tabview.add("📑 Monthly Report")

# ================= TAB 0: ERP DASHBOARD =================
dash_top = ctk.CTkFrame(tab_dash, fg_color="transparent")
dash_top.pack(fill="x", padx=15, pady=(8, 4))

ctk.CTkLabel(dash_top, text="School ERP Overview & Live Status", font=ctk.CTkFont(size=20, weight="bold")).pack(side="left")

frame_dash_date = ctk.CTkFrame(dash_top, fg_color="transparent")
frame_dash_date.pack(side="right")

ctk.CTkLabel(frame_dash_date, text="तारीख चुनें:", font=ctk.CTkFont(weight="bold")).pack(side="left", padx=4)
entry_dash_date = ctk.CTkEntry(frame_dash_date, width=110)
entry_dash_date.insert(0, datetime.today().strftime("%d-%m-%Y"))
entry_dash_date.pack(side="left", padx=4)

ctk.CTkButton(frame_dash_date, text="📅", width=35, command=lambda: open_calendar(entry_dash_date, refresh_dashboard)).pack(side="left", padx=2)
ctk.CTkButton(frame_dash_date, text="🔄 लोड करें", width=80, command=refresh_dashboard).pack(side="left", padx=4)

cards_frame = ctk.CTkFrame(tab_dash, fg_color="transparent")
cards_frame.pack(fill="x", padx=15, pady=8)
cards_frame.grid_columnconfigure((0, 1, 2, 3, 4), weight=1, uniform="card")

def create_kpi_card(parent, col, title, initial_val, color_accent):
    card = ctk.CTkFrame(parent, corner_radius=10, border_width=2, border_color=color_accent)
    card.grid(row=0, column=col, padx=6, pady=4, sticky="nsew")
    ctk.CTkLabel(card, text=title, font=ctk.CTkFont(size=13)).pack(pady=(8, 2))
    val_lbl = ctk.CTkLabel(card, text=initial_val, font=ctk.CTkFont(size=24, weight="bold"), text_color=color_accent)
    val_lbl.pack(pady=(0, 8))
    return val_lbl

lbl_val_students = create_kpi_card(cards_frame, 0, "कुल विद्यार्थी", "0", "#1976d2")
lbl_val_present = create_kpi_card(cards_frame, 1, "उस दिन उपस्थित (P)", "0", "#2e7d32")
lbl_val_absent = create_kpi_card(cards_frame, 2, "उस दिन अनुपस्थित (A)", "0", "#c62828")
lbl_val_leave = create_kpi_card(cards_frame, 3, "छुट्टी / अवकाश", "0", "#ef6c00")
lbl_val_routes = create_kpi_card(cards_frame, 4, "सक्रिय बस रूट्स", "0", "#6a1b9a")

nav_frame = ctk.CTkFrame(tab_dash, fg_color=("gray90", "gray20"), height=42, corner_radius=8)
nav_frame.pack(fill="x", padx=15, pady=4)

ctk.CTkLabel(nav_frame, text="⚡ Quick Access:", font=ctk.CTkFont(size=13, weight="bold")).pack(side="left", padx=12)
ctk.CTkButton(nav_frame, text="➕ नया एडमिशन", width=130, height=28, command=lambda: tabview.set("➕ New Admission")).pack(side="left", padx=5)
ctk.CTkButton(nav_frame, text="📋 रिकॉर्ड्स लिस्ट", width=130, height=28, command=lambda: tabview.set("📋 Student Records")).pack(side="left", padx=5)
ctk.CTkButton(nav_frame, text="📅 हाजिरी लगाएं / देखें", width=150, height=28, fg_color="#2e7d32", hover_color="#1b5e20", command=lambda: tabview.set("📅 Daily Attendance")).pack(side="left", padx=5)
ctk.CTkButton(nav_frame, text="📑 मासिक रजिस्टर", width=140, height=28, fg_color="#6a1b9a", hover_color="#4a148c", command=lambda: tabview.set("📑 Monthly Report")).pack(side="left", padx=5)

ctk.CTkLabel(tab_dash, text="📋 कक्षा-वार उपस्थिति रिपोर्ट (क्रम: PP.3+, PP.4+, PP.5+, 1, 2, 3... 10):", font=ctk.CTkFont(size=13, weight="bold")).pack(anchor="w", padx=20, pady=(8, 3))

dash_table_frame = ctk.CTkFrame(tab_dash)
dash_table_frame.pack(fill="both", expand=True, padx=15, pady=(0, 12))

dash_cols = ("Class", "Total Students", "Present", "Absent", "Leave/Holiday", "Attendance %", "Status")
tree_dash = ttk.Treeview(dash_table_frame, columns=dash_cols, show="headings", selectmode="browse")

dash_scroll_y = ttk.Scrollbar(dash_table_frame, orient="vertical", command=tree_dash.yview)
tree_dash.configure(yscrollcommand=dash_scroll_y.set)
dash_scroll_y.pack(side="right", fill="y")
tree_dash.pack(fill="both", expand=True)

for col in dash_cols:
    tree_dash.heading(col, text=col)
    tree_dash.column(col, width=120, anchor="center")

# ================= TAB 1: ENTRY FORM =================
lbl_form_title = ctk.CTkLabel(tab_entry, text="विद्यार्थी प्रवेश फ़ॉर्म (Student Admission Form)", font=ctk.CTkFont(size=18, weight="bold"))
lbl_form_title.pack(pady=5)

form_frame = ctk.CTkScrollableFrame(tab_entry, fg_color="transparent")
form_frame.pack(fill="both", expand=True, padx=10, pady=5)
form_frame.grid_columnconfigure(0, weight=1)
form_frame.grid_columnconfigure(1, weight=1)

ctk.CTkLabel(form_frame, text="SR No.*:").grid(row=0, column=0, padx=20, pady=(5, 0), sticky="w")
entry_sr = ctk.CTkEntry(form_frame, placeholder_text="उदा. 1054", width=330)
entry_sr.grid(row=1, column=0, padx=20, pady=(0, 8), sticky="w")

ctk.CTkLabel(form_frame, text="Date of Admission:").grid(row=0, column=1, padx=20, pady=(5, 0), sticky="w")
frame_doa = ctk.CTkFrame(form_frame, fg_color="transparent")
frame_doa.grid(row=1, column=1, padx=20, pady=(0, 8), sticky="w")
entry_doa = ctk.CTkEntry(frame_doa, placeholder_text="DD-MM-YYYY", width=270)
entry_doa.pack(side="left", padx=(0, 5))
ctk.CTkButton(frame_doa, text="📅", width=40, command=lambda: open_calendar(entry_doa)).pack(side="left")

ctk.CTkLabel(form_frame, text="Student Name*:").grid(row=2, column=0, padx=20, pady=(5, 0), sticky="w")
entry_name = ctk.CTkEntry(form_frame, placeholder_text="विद्यार्थी का नाम", width=330)
entry_name.grid(row=3, column=0, padx=20, pady=(0, 8), sticky="w")

ctk.CTkLabel(form_frame, text="Class*:").grid(row=2, column=1, padx=20, pady=(5, 0), sticky="w")
entry_class = ctk.CTkEntry(form_frame, placeholder_text="उदा. PP.3+ / 1 / 6", width=330)
entry_class.grid(row=3, column=1, padx=20, pady=(0, 8), sticky="w")

ctk.CTkLabel(form_frame, text="Father Name:").grid(row=4, column=0, padx=20, pady=(5, 0), sticky="w")
entry_father = ctk.CTkEntry(form_frame, placeholder_text="पिता का नाम", width=330)
entry_father.grid(row=5, column=0, padx=20, pady=(0, 8), sticky="w")

ctk.CTkLabel(form_frame, text="Mother Name:").grid(row=4, column=1, padx=20, pady=(5, 0), sticky="w")
entry_mother = ctk.CTkEntry(form_frame, placeholder_text="माता का नाम", width=330)
entry_mother.grid(row=5, column=1, padx=20, pady=(0, 8), sticky="w")

ctk.CTkLabel(form_frame, text="Date of Birth (DOB):").grid(row=6, column=0, padx=20, pady=(5, 0), sticky="w")
frame_dob = ctk.CTkFrame(form_frame, fg_color="transparent")
frame_dob.grid(row=7, column=0, padx=20, pady=(0, 8), sticky="w")
entry_dob = ctk.CTkEntry(frame_dob, placeholder_text="DD-MM-YYYY", width=270)
entry_dob.pack(side="left", padx=(0, 5))
ctk.CTkButton(frame_dob, text="📅", width=40, command=lambda: open_calendar(entry_dob)).pack(side="left")

ctk.CTkLabel(form_frame, text="Gender:").grid(row=6, column=1, padx=20, pady=(5, 0), sticky="w")
opt_gender = ctk.CTkOptionMenu(form_frame, values=["Male", "Female", "Other"], width=330)
opt_gender.grid(row=7, column=1, padx=20, pady=(0, 8), sticky="w")

ctk.CTkLabel(form_frame, text="Social Category:").grid(row=8, column=0, padx=20, pady=(5, 0), sticky="w")
opt_category = ctk.CTkOptionMenu(form_frame, values=["General", "OBC", "SC", "ST", "SBC/MBC", "EWS"], width=330)
opt_category.grid(row=9, column=0, padx=20, pady=(0, 10), sticky="w")

ctk.CTkLabel(form_frame, text="RTE:").grid(row=8, column=1, padx=20, pady=(5, 0), sticky="w")
opt_rte = ctk.CTkOptionMenu(form_frame, values=["No", "Yes"], width=330)
opt_rte.grid(row=9, column=1, padx=20, pady=(0, 10), sticky="w")

ctk.CTkLabel(form_frame, text="Mobile Number:").grid(row=10, column=0, padx=20, pady=(5, 0), sticky="w")
entry_mobile = ctk.CTkEntry(form_frame, placeholder_text="10 अंकों का मोबाइल नंबर", width=330)
entry_mobile.grid(row=11, column=0, padx=20, pady=(0, 8), sticky="w")

ctk.CTkLabel(form_frame, text="Route:").grid(row=10, column=1, padx=20, pady=(5, 0), sticky="w")
opt_route = ctk.CTkOptionMenu(form_frame, values=route_options, width=330)
opt_route.set("Self")
opt_route.grid(row=11, column=1, padx=20, pady=(0, 8), sticky="w")

ctk.CTkLabel(form_frame, text="Address:").grid(row=12, column=0, padx=20, pady=(5, 0), sticky="w")
entry_address = ctk.CTkEntry(form_frame, placeholder_text="गाँव / ढाणी / शहर का पता", width=710)
entry_address.grid(row=13, column=0, columnspan=2, padx=20, pady=(0, 10), sticky="w")

# बटन एरिया: Save, Import Excel, और Clear Form
btn_frame = ctk.CTkFrame(tab_entry, fg_color="transparent")
btn_frame.pack(pady=10)

btn_save = ctk.CTkButton(btn_frame, text="Save Student", width=140, height=36, command=save_or_update)
btn_save.pack(side="left", padx=8)

btn_import_excel = ctk.CTkButton(btn_frame, text="📁 Import Excel", width=140, height=36, fg_color="#2e7d32", hover_color="#1b5e20", command=import_excel)
btn_import_excel.pack(side="left", padx=8)

btn_clear = ctk.CTkButton(btn_frame, text="Clear / Cancel", width=110, height=36, fg_color="gray", command=clear_fields)
btn_clear.pack(side="left", padx=8)

# ================= TAB 2: RECORDS TABLE =================
search_frame = ctk.CTkFrame(tab_records, fg_color="transparent")
search_frame.pack(fill="x", padx=15, pady=8)

entry_search = ctk.CTkEntry(search_frame, placeholder_text="SR No., नाम, कक्षा या रूट से खोजें...", width=320)
entry_search.pack(side="left", padx=5)

ctk.CTkButton(search_frame, text="🔍 Search", width=90, command=search_data).pack(side="left", padx=5)
ctk.CTkButton(search_frame, text="Reset", width=70, fg_color="gray", command=reset_search).pack(side="left", padx=5)

lbl_count = ctk.CTkLabel(search_frame, text="कुल रिकॉर्ड: 0", font=ctk.CTkFont(weight="bold"))
lbl_count.pack(side="left", padx=20)

ctk.CTkButton(search_frame, text="📊 Export to Excel", fg_color="#2e7d32", hover_color="#1b5e20", width=140, command=export_excel).pack(side="right", padx=5)
ctk.CTkButton(search_frame, text="🗑️ Delete Selected", fg_color="#c62828", hover_color="#8e0000", width=130, command=delete_selected).pack(side="right", padx=5)

info_lbl = ctk.CTkLabel(tab_records, text="💡 टिप: किसी भी छात्र का डेटा सुधारने (Edit) के लिए उसकी लाइन पर Double-Click करें।", font=ctk.CTkFont(size=12), text_color="gray")
info_lbl.pack(anchor="w", padx=20, pady=(0, 5))

table_frame = ctk.CTkFrame(tab_records)
table_frame.pack(fill="both", expand=True, padx=15, pady=(0, 10))

columns = ("SR No", "Name", "Class", "Father Name", "Mobile", "Route", "Category", "RTE", "DOB", "DOA", "Address")
tree = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="browse")

scroll_y = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
scroll_x = ttk.Scrollbar(table_frame, orient="horizontal", command=tree.xview)
tree.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)

scroll_y.pack(side="right", fill="y")
scroll_x.pack(side="bottom", fill="x")
tree.pack(fill="both", expand=True)

col_widths = {
    "SR No": 70, "Name": 130, "Class": 80, "Father Name": 130, 
    "Mobile": 100, "Route": 110, "Category": 80, "RTE": 50, 
    "DOB": 90, "DOA": 90, "Address": 150
}
for col in columns:
    tree.heading(col, text=col)
    tree.column(col, width=col_widths.get(col, 100), anchor="center" if col in ["SR No", "Class", "RTE", "DOB", "DOA"] else "w")

tree.bind("<Double-1>", on_double_click_row)

# ================= TAB 3: DAILY ATTENDANCE =================
att_top_frame = ctk.CTkFrame(tab_attendance, fg_color="transparent")
att_top_frame.pack(fill="x", padx=15, pady=(8, 2))

ctk.CTkLabel(att_top_frame, text="तारीख (Back-Date/आज):", font=ctk.CTkFont(weight="bold")).pack(side="left", padx=(5, 2))
today_str = datetime.today().strftime("%d-%m-%Y")
entry_att_date = ctk.CTkEntry(att_top_frame, width=110)
entry_att_date.insert(0, today_str)
entry_att_date.pack(side="left", padx=(0, 5))
ctk.CTkButton(att_top_frame, text="📅", width=35, command=lambda: open_calendar(entry_att_date)).pack(side="left", padx=(0, 15))

ctk.CTkLabel(att_top_frame, text="कक्षा चुनें:", font=ctk.CTkFont(weight="bold")).pack(side="left", padx=(5, 2))
opt_att_class = ctk.CTkOptionMenu(att_top_frame, values=["लोड हो रहा है..."], width=130)
opt_att_class.pack(side="left", padx=(0, 15))

ctk.CTkButton(att_top_frame, text="📥 Load / View", width=120, command=load_attendance_list).pack(side="left", padx=5)

lbl_att_summary = ctk.CTkLabel(att_top_frame, text="", font=ctk.CTkFont(weight="bold"))
lbl_att_summary.pack(side="left", padx=15)

ctk.CTkButton(att_top_frame, text="💾 Save / Update", fg_color="#2e7d32", hover_color="#1b5e20", width=150, command=save_attendance).pack(side="right", padx=5)

lbl_att_status_mode = ctk.CTkLabel(tab_attendance, text="", font=ctk.CTkFont(size=12, weight="bold"))
lbl_att_status_mode.pack(anchor="w", padx=20, pady=(2, 4))

quick_action_frame = ctk.CTkFrame(tab_attendance, fg_color=("gray90", "gray20"), height=38, corner_radius=6)
quick_action_frame.pack(fill="x", padx=15, pady=(2, 6))

ctk.CTkLabel(quick_action_frame, text="⚡ Quick Actions (1-Click):", font=ctk.CTkFont(size=12, weight="bold")).pack(side="left", padx=12)

ctk.CTkButton(quick_action_frame, text="🟢 All Present", fg_color="#2e7d32", hover_color="#1b5e20", width=110, height=26, command=lambda: set_all_status("Present")).pack(side="left", padx=5, pady=4)
ctk.CTkButton(quick_action_frame, text="🔴 All Absent", fg_color="#c62828", hover_color="#8e0000", width=110, height=26, command=lambda: set_all_status("Absent")).pack(side="left", padx=5, pady=4)
ctk.CTkButton(quick_action_frame, text="🟡 All Leave", fg_color="#ef6c00", hover_color="#b24d00", width=105, height=26, command=lambda: set_all_status("Leave")).pack(side="left", padx=5, pady=4)
ctk.CTkButton(quick_action_frame, text="🏖️ Mark Holiday", fg_color="#6a1b9a", hover_color="#4a148c", width=125, height=26, command=lambda: set_all_status("Holiday")).pack(side="left", padx=5, pady=4)

att_scroll_frame = ctk.CTkScrollableFrame(tab_attendance)
att_scroll_frame.pack(fill="both", expand=True, padx=15, pady=(0, 10))

# ================= TAB 4: ATTENDANCE REPORT =================
rep_top_frame = ctk.CTkFrame(tab_report, fg_color="transparent")
rep_top_frame.pack(fill="x", padx=15, pady=10)

ctk.CTkLabel(rep_top_frame, text="कक्षा:").pack(side="left", padx=(5, 2))
opt_rep_class = ctk.CTkOptionMenu(rep_top_frame, values=["लोड हो रहा है..."], width=110)
opt_rep_class.pack(side="left", padx=(0, 10))

ctk.CTkLabel(rep_top_frame, text="महीना:").pack(side="left", padx=(5, 2))
opt_rep_month = ctk.CTkOptionMenu(rep_top_frame, values=months_list, width=110)
curr_m = datetime.today().month
opt_rep_month.set(months_list[curr_m - 1])
opt_rep_month.pack(side="left", padx=(0, 10))

ctk.CTkLabel(rep_top_frame, text="साल:").pack(side="left", padx=(5, 2))
opt_rep_year = ctk.CTkOptionMenu(rep_top_frame, values=years_list, width=90)
opt_rep_year.set(str(datetime.today().year))
opt_rep_year.pack(side="left", padx=(0, 15))

ctk.CTkButton(rep_top_frame, text="📊 View Report", width=120, command=load_attendance_report).pack(side="left", padx=5)
ctk.CTkButton(rep_top_frame, text="📥 Export Monthly Sheet", fg_color="#2e7d32", hover_color="#1b5e20", width=170, command=export_attendance_monthly).pack(side="right", padx=5)

lbl_rep_summary = ctk.CTkLabel(tab_report, text="", font=ctk.CTkFont(size=12, weight="bold"), text_color="gray")
lbl_rep_summary.pack(anchor="w", padx=20, pady=(0, 8))

rep_table_frame = ctk.CTkFrame(tab_report)
rep_table_frame.pack(fill="both", expand=True, padx=15, pady=(0, 15))

rep_columns = ("SR No", "Name", "Father Name", "Class", "Working Days", "Present (P)", "Absent (A)", "Leave (L)", "Holiday (H)", "Percentage")
tree_report = ttk.Treeview(rep_table_frame, columns=rep_columns, show="headings", selectmode="browse")

rep_scroll_y = ttk.Scrollbar(rep_table_frame, orient="vertical", command=tree_report.yview)
rep_scroll_x = ttk.Scrollbar(rep_table_frame, orient="horizontal", command=tree_report.xview)
tree_report.configure(yscrollcommand=rep_scroll_y.set, xscrollcommand=rep_scroll_x.set)

rep_scroll_y.pack(side="right", fill="y")
rep_scroll_x.pack(side="bottom", fill="x")
tree_report.pack(fill="both", expand=True)

rep_col_widths = {
    "SR No": 75, "Name": 130, "Father Name": 130, "Class": 75, 
    "Working Days": 95, "Present (P)": 85, "Absent (A)": 85, "Leave (L)": 75, "Holiday (H)": 85, "Percentage": 85
}
for col in rep_columns:
    tree_report.heading(col, text=col)
    tree_report.column(col, width=rep_col_widths.get(col, 100), anchor="center" if col not in ["Name", "Father Name"] else "w")

load_records()
update_class_dropdown()
refresh_dashboard()

app.mainloop()